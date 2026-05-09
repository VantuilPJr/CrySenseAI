from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle, PageBreak


ROOT = Path(__file__).resolve().parents[1]
SOURCE_MD = ROOT / "DOCUMENTACAO_TDE_PRE_BANCA_ABNT.md"
OUT_DOCX = ROOT / "DOCUMENTACAO_TDE_PRE_BANCA_ABNT_FINAL.docx"
OUT_PDF = ROOT / "DOCUMENTACAO_TDE_PRE_BANCA_ABNT_FINAL.pdf"


@dataclass
class Block:
    kind: str
    level: int | None
    text: str | None
    rows: list[list[str]] | None = None


def parse_markdown(text: str) -> list[Block]:
    lines = text.splitlines()
    blocks: list[Block] = []
    paragraph_buffer: list[str] = []
    i = 0

    def flush_paragraph() -> None:
        nonlocal paragraph_buffer
        if paragraph_buffer:
            joined = " ".join(part.strip() for part in paragraph_buffer if part.strip())
            if joined:
                blocks.append(Block("paragraph", None, joined))
            paragraph_buffer = []

    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()

        if not stripped:
            flush_paragraph()
            i += 1
            continue

        if stripped.startswith("## "):
            flush_paragraph()
            blocks.append(Block("heading", 2, stripped[3:].strip()))
            i += 1
            continue

        if stripped.startswith("### "):
            flush_paragraph()
            blocks.append(Block("heading", 3, stripped[4:].strip()))
            i += 1
            continue

        if stripped.startswith("|"):
            flush_paragraph()
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                row = [cell.strip() for cell in lines[i].strip().strip("|").split("|")]
                table_lines.append(row)
                i += 1
            if len(table_lines) >= 2 and all(re.fullmatch(r"-+", cell.replace(":", "").replace(" ", "")) for cell in table_lines[1]):
                table_lines.pop(1)
            blocks.append(Block("table", None, None, table_lines))
            continue

        if re.match(r"^- \[[ xX]\] ", stripped):
            flush_paragraph()
            blocks.append(Block("bullet", None, stripped[6:].strip()))
            i += 1
            continue

        if stripped.startswith("- "):
            flush_paragraph()
            blocks.append(Block("bullet", None, stripped[2:].strip()))
            i += 1
            continue

        if re.match(r"^\d+\. ", stripped):
            flush_paragraph()
            blocks.append(Block("numbered", None, stripped))
            i += 1
            continue

        paragraph_buffer.append(stripped.replace("  ", " "))
        i += 1

    flush_paragraph()
    return blocks


def sanitize_inline(text: str) -> str:
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\([^\)]+\)", r"\1", text)
    return text.strip()


def apply_doc_defaults(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(3)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2)
    section.bottom_margin = Cm(2)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Cm(1.25)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.space_before = Pt(0)

    if "ABNT Heading 1" not in document.styles:
        style = document.styles.add_style("ABNT Heading 1", WD_STYLE_TYPE.PARAGRAPH)
    else:
        style = document.styles["ABNT Heading 1"]
    style.base_style = document.styles["Normal"]
    style.font.bold = True
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    style.font.size = Pt(12)
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    style.paragraph_format.first_line_indent = Cm(0)
    style.paragraph_format.space_before = Pt(12)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.5

    if "ABNT Heading 2" not in document.styles:
        style2 = document.styles.add_style("ABNT Heading 2", WD_STYLE_TYPE.PARAGRAPH)
    else:
        style2 = document.styles["ABNT Heading 2"]
    style2.base_style = document.styles["Normal"]
    style2.font.bold = True
    style2.font.name = "Times New Roman"
    style2._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    style2.font.size = Pt(12)
    style2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    style2.paragraph_format.first_line_indent = Cm(0)
    style2.paragraph_format.space_before = Pt(10)
    style2.paragraph_format.space_after = Pt(4)
    style2.paragraph_format.line_spacing = 1.5


def add_cover_page(document: Document, blocks: list[Block], start_idx: int) -> int:
    idx = start_idx + 1
    lines: list[str] = []
    while idx < len(blocks) and not (blocks[idx].kind == "heading" and blocks[idx].level == 2):
        if blocks[idx].kind == "paragraph":
            lines.append(sanitize_inline(blocks[idx].text or ""))
        idx += 1

    for pos, line in enumerate(lines):
        para = document.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.first_line_indent = Cm(0)
        para.paragraph_format.line_spacing = 1.5
        if pos in {0, 1, 2, len(lines) - 2, len(lines) - 1}:
            para.paragraph_format.space_before = Pt(0)
        elif pos == 3:
            para.paragraph_format.space_before = Pt(72)
        elif pos == len(lines) - 2:
            para.paragraph_format.space_before = Pt(120)
        run = para.add_run(line)
        run.font.name = "Times New Roman"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        run.font.size = Pt(12)
        if pos == 4:
            run.bold = True

    document.add_page_break()
    return idx


def add_title_page(document: Document, blocks: list[Block], start_idx: int) -> int:
    idx = start_idx + 1
    lines: list[str] = []
    while idx < len(blocks) and not (blocks[idx].kind == "heading" and blocks[idx].level == 2):
        if blocks[idx].kind == "paragraph":
            lines.append(sanitize_inline(blocks[idx].text or ""))
        idx += 1

    for pos, line in enumerate(lines):
        para = document.add_paragraph()
        para.paragraph_format.first_line_indent = Cm(0)
        para.paragraph_format.line_spacing = 1.5
        if pos in {0, 1, len(lines) - 2, len(lines) - 1}:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif pos == 2:
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            para.paragraph_format.left_indent = Cm(8)
        else:
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if pos == 1:
            para.paragraph_format.space_before = Pt(72)
        if pos == 2:
            para.paragraph_format.space_before = Pt(72)
        if pos == len(lines) - 2:
            para.paragraph_format.space_before = Pt(100)
        run = para.add_run(line)
        run.font.name = "Times New Roman"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        run.font.size = Pt(12)
        if pos == 1:
            run.bold = True

    document.add_page_break()
    return idx


def style_table(table) -> None:
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.first_line_indent = Cm(0)
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.15
                if row_idx == 0:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
                    run.font.size = Pt(10)
                    if row_idx == 0:
                        run.bold = True


def set_cell_text(cell, text: str) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.add_run(sanitize_inline(text))


def generate_docx(blocks: list[Block]) -> None:
    document = Document()
    apply_doc_defaults(document)

    idx = 0
    while idx < len(blocks):
        block = blocks[idx]
        if block.kind == "heading" and block.level == 2 and block.text == "CAPA":
            idx = add_cover_page(document, blocks, idx)
            continue
        if block.kind == "heading" and block.level == 2 and block.text == "FOLHA DE ROSTO":
            idx = add_title_page(document, blocks, idx)
            continue

        if block.kind == "heading":
            para = document.add_paragraph(style="ABNT Heading 1" if block.level == 2 else "ABNT Heading 2")
            para.add_run(sanitize_inline(block.text or ""))
        elif block.kind == "paragraph":
            para = document.add_paragraph(style="Normal")
            para.add_run(sanitize_inline(block.text or ""))
            if block.text and (block.text.startswith("Palavras-chave:") or block.text.startswith("Keywords:")):
                para.paragraph_format.first_line_indent = Cm(0)
        elif block.kind == "bullet":
            para = document.add_paragraph(style="Normal")
            para.paragraph_format.left_indent = Cm(1.25)
            para.paragraph_format.first_line_indent = Cm(0)
            para.add_run("• " + sanitize_inline(block.text or ""))
        elif block.kind == "numbered":
            para = document.add_paragraph(style="Normal")
            para.paragraph_format.left_indent = Cm(1.25)
            para.paragraph_format.first_line_indent = Cm(0)
            para.add_run(sanitize_inline(block.text or ""))
        elif block.kind == "table" and block.rows:
            table = document.add_table(rows=len(block.rows), cols=len(block.rows[0]))
            for row_idx, row in enumerate(block.rows):
                for col_idx, value in enumerate(row):
                    set_cell_text(table.cell(row_idx, col_idx), value)
            style_table(table)

        idx += 1

    document.save(OUT_DOCX)


def xml_escape(text: str) -> str:
    return (
        sanitize_inline(text)
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )


def build_pdf_styles():
    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="ABNTNormal",
            fontName="Times-Roman",
            fontSize=12,
            leading=18,
            alignment=TA_JUSTIFY,
            firstLineIndent=1.25 * cm,
            spaceAfter=0,
            spaceBefore=0,
        )
    )
    styles.add(
        ParagraphStyle(
            name="ABNTHeading1",
            fontName="Times-Bold",
            fontSize=12,
            leading=18,
            alignment=TA_LEFT,
            spaceBefore=12,
            spaceAfter=6,
        )
    )
    styles.add(
        ParagraphStyle(
            name="ABNTHeading2",
            fontName="Times-Bold",
            fontSize=12,
            leading=18,
            alignment=TA_LEFT,
            spaceBefore=10,
            spaceAfter=4,
        )
    )
    styles.add(
        ParagraphStyle(
            name="ABNTCenter",
            fontName="Times-Roman",
            fontSize=12,
            leading=18,
            alignment=TA_CENTER,
        )
    )
    styles.add(
        ParagraphStyle(
            name="ABNTBullet",
            fontName="Times-Roman",
            fontSize=12,
            leading=18,
            leftIndent=1.25 * cm,
            firstLineIndent=0,
            alignment=TA_LEFT,
        )
    )
    styles.add(
        ParagraphStyle(
            name="ABNTKeywords",
            fontName="Times-Roman",
            fontSize=12,
            leading=18,
            firstLineIndent=0,
            alignment=TA_JUSTIFY,
        )
    )
    return styles


def cover_story(lines: list[str], styles) -> list:
    story = [Spacer(1, 2 * cm)]
    for pos, line in enumerate(lines):
        story.append(Paragraph(xml_escape(line), styles["ABNTCenter"]))
        if pos == 2:
            story.append(Spacer(1, 2.5 * cm))
        elif pos == 4:
            story.append(Spacer(1, 5 * cm))
        else:
            story.append(Spacer(1, 0.3 * cm))
    story.append(PageBreak())
    return story


def title_page_story(lines: list[str], styles) -> list:
    story = [Spacer(1, 2 * cm)]
    for pos, line in enumerate(lines):
        if pos in {0, 1, len(lines) - 2, len(lines) - 1}:
            style = styles["ABNTCenter"]
        elif pos == 2:
            shifted = ParagraphStyle(
                name="TempTitlePage",
                parent=styles["ABNTNormal"],
                leftIndent=8 * cm,
                firstLineIndent=0,
            )
            style = shifted
        else:
            style = styles["ABNTNormal"]
        story.append(Paragraph(xml_escape(line), style))
        if pos == 1:
            story.append(Spacer(1, 2.5 * cm))
        elif pos == 2:
            story.append(Spacer(1, 4 * cm))
        else:
            story.append(Spacer(1, 0.3 * cm))
    story.append(PageBreak())
    return story


def generate_pdf(blocks: list[Block]) -> None:
    styles = build_pdf_styles()
    story = []
    idx = 0
    while idx < len(blocks):
        block = blocks[idx]
        if block.kind == "heading" and block.level == 2 and block.text == "CAPA":
            idx += 1
            lines = []
            while idx < len(blocks) and not (blocks[idx].kind == "heading" and blocks[idx].level == 2):
                if blocks[idx].kind == "paragraph":
                    lines.append(blocks[idx].text or "")
                idx += 1
            story.extend(cover_story(lines, styles))
            continue
        if block.kind == "heading" and block.level == 2 and block.text == "FOLHA DE ROSTO":
            idx += 1
            lines = []
            while idx < len(blocks) and not (blocks[idx].kind == "heading" and blocks[idx].level == 2):
                if blocks[idx].kind == "paragraph":
                    lines.append(blocks[idx].text or "")
                idx += 1
            story.extend(title_page_story(lines, styles))
            continue

        if block.kind == "heading":
            style = styles["ABNTHeading1"] if block.level == 2 else styles["ABNTHeading2"]
            story.append(Paragraph(xml_escape(block.text or ""), style))
        elif block.kind == "paragraph":
            style = styles["ABNTKeywords"] if (block.text or "").startswith(("Palavras-chave:", "Keywords:")) else styles["ABNTNormal"]
            story.append(Paragraph(xml_escape(block.text or ""), style))
        elif block.kind == "bullet":
            story.append(Paragraph(xml_escape("• " + (block.text or "")), styles["ABNTBullet"]))
        elif block.kind == "numbered":
            story.append(Paragraph(xml_escape(block.text or ""), styles["ABNTBullet"]))
        elif block.kind == "table" and block.rows:
            data = [[Paragraph(xml_escape(cell), styles["ABNTBullet"]) for cell in row] for row in block.rows]
            table = Table(data, repeatRows=1)
            table.setStyle(
                TableStyle(
                    [
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EDEDED")),
                        ("FONTNAME", (0, 0), (-1, 0), "Times-Bold"),
                        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 4),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                        ("TOPPADDING", (0, 0), (-1, -1), 4),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                    ]
                )
            )
            story.append(table)
        story.append(Spacer(1, 0.15 * cm))
        idx += 1

    doc = SimpleDocTemplate(
        str(OUT_PDF),
        pagesize=A4,
        leftMargin=3 * cm,
        rightMargin=2 * cm,
        topMargin=3 * cm,
        bottomMargin=2 * cm,
        title="CrySense AI v2.0 - Documentacao Pre-Banca ABNT",
    )
    doc.build(story)


def main() -> None:
    content = SOURCE_MD.read_text(encoding="utf-8")
    blocks = parse_markdown(content)
    generate_docx(blocks)
    try:
        generate_pdf(blocks)
        pdf_message = f"PDF generated: {OUT_PDF}"
    except PermissionError:
        pdf_message = f"PDF not updated because the file is open/locked: {OUT_PDF}"
    print(f"DOCX generated: {OUT_DOCX}")
    print(pdf_message)


if __name__ == "__main__":
    main()